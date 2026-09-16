import os

from docx import Document
from pypdf import PdfReader


def _table_to_text(table) -> str:
    """Read a docx table as one line per row, dropping repeated text from merged cells."""
    lines = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            text = cell.text.strip()
            # A horizontally merged cell repeats its text for every column it spans
            if text and (not cells or text != cells[-1]):
                cells.append(text)
        if cells:
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def extract_text(file_path: str) -> str:
    """Extract text from a .docx or .pdf file, chosen by its extension."""
    if not os.path.isfile(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".docx":
        document = Document(file_path)
        paragraph_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        table_text = "\n".join(_table_to_text(table) for table in document.tables)
        return "\n".join(part for part in (paragraph_text, table_text) if part)

    if ext == ".pdf":
        reader = PdfReader(file_path)
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    raise ValueError(f"Unsupported file type: {ext}")
